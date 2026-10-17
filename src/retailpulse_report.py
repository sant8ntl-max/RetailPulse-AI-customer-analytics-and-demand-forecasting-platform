"""
RetailPulse - automatic project report (DOCX) generator.
Reads outputs/metrics.json + outputs/charts_index.json produced by the notebooks, so every number
and every chart in the report comes from the real run (nothing is hard-coded).

    import retailpulse_report as rr
    rr.build_report("RetailPulse_Project_Report_FULL.docx", compact=False)   # every chart, long
    rr.build_report("RetailPulse_Zidio_Submission_Report.docx", compact=True)  # ~15 pages for the Zidio PDF

Author: Santosh Kumar
"""
from __future__ import annotations

import io
import json
import subprocess
import sys
from pathlib import Path

try:
    import docx  # noqa: F401
except ImportError:  # Colab: install on demand
    subprocess.run([sys.executable, "-m", "pip", "install", "-q", "python-docx"], check=False)

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import retailpulse_utils as ru

AUTHOR = "Santosh Kumar"
NAVY, TEAL, GREY = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x0F, 0x76, 0x6E), RGBColor(0x64, 0x74, 0x8B)

SECTION_TEXT = {
    "EDA": ("Exploratory Data Analysis (Notebook 01)",
            "EDA was the first step: understand the raw table, its quality problems, the seasonality of demand, where revenue comes from "
            "(products, categories, countries, customers) and how variables relate. Every later model depends on these findings."),
    "Cleaning & Feature Engineering": ("Data Cleaning, Validation & Feature Engineering (Notebook 02)",
            "Cancellations, invalid prices, non-product codes, duplicates and extreme outliers were removed with a documented waterfall. "
            "Ten validation expectations (Great-Expectations style) were checked. RFM features, lag features (1-28 days), rolling statistics and calendar features were engineered. "
            "Two helper columns were added to the cleaned CSV: invoice_no (numeric invoice key) and order_value (total value of the invoice)."),
    "Customer Segmentation": ("Customer Segmentation (Notebook 03)",
            "Customers were segmented with RFM scoring (rule-based), K-Means (K chosen with silhouette inside the business range 6-8) and DBSCAN (outlier discovery). "
            "Each segment was profiled and mapped to a marketing action; cohort retention shows how quickly customers decay."),
    "Time-Series Preparation": ("Time-Series Preparation (Notebook 04)",
            "A continuous daily demand series was built, closed weekdays detected, the series decomposed (classical + STL), stationarity tested (ADF / KPSS), "
            "ACF/PACF inspected, outliers flagged and a leakage-free train / validation / test split of 30 days each created."),
    "Forecasting - Prophet": ("Baseline Forecasting with Prophet (Notebook 05)",
            "Prophet with weekly + yearly seasonality and UK holidays was fitted as the baseline, compared with Seasonal Naive and Holt-Winters, diagnosed with residual plots and tested for hyper-parameter sensitivity."),
    "Forecasting - LSTM": ("Deep-Learning Forecasting with LSTM (Notebook 06)",
            "A PyTorch LSTM uses a 28-day window plus known-future calendar covariates and produces a recursive 30-day forecast with early stopping."),
    "MLOps - MLflow": ("Experiment Tracking with MLflow (Notebooks 07 and 14)",
            "All baseline and advanced runs were logged (parameters, metrics, artifacts) in a local SQLite-backed MLflow store, compared against the acceptance targets and summarised in a scorecard."),
    "Forecasting - Hybrid": ("Hybrid Prophet + LSTM Ensemble (Notebook 08)",
            "The blend weight was optimised on the validation window; models were re-fitted on train + validation to score the test window and finally refit on all data for the 30-day production forecast."),
    "Churn Prediction": ("Churn Prediction with XGBoost + SHAP (Notebook 09)",
            "A leakage-free churn label (no purchase in the next 90 days) was modelled with Logistic Regression, Random Forest and XGBoost. Performance, calibration, gain/lift and SHAP explanations are shown."),
    "Inventory Optimization": ("Inventory Optimization (Notebook 10)",
            "Safety stock, reorder point, EOQ and order-up-to level were computed per SKU from forecast demand and a 30-day policy back-test compared a naive buffer rule with the forecast-driven policy."),
    "Tuning": ("Hyper-parameter Tuning with Optuna (Notebook 11)",
            "Bayesian (TPE) optimisation tuned XGBoost (cross-validated AUC) and Prophet (validation WAPE); results were compared with the defaults on the untouched test set."),
    "Drift Detection": ("Drift Detection & Monitoring (Notebook 12)",
            "PSI and KS tests compare a reference and a current period; an injected price / demand shock demonstrates drift detection and the automatic retrain trigger."),
    "Model Refinement": ("Model Refinement & Retraining Pipeline (Notebook 13)",
            "Calibration, profit-optimal thresholding, feature-group ablation, ensembling, learning curves and weekday bias correction were evaluated; an Airflow DAG automates retraining."),
}
SECTION_ORDER = ["EDA", "Cleaning & Feature Engineering", "Customer Segmentation", "Time-Series Preparation", "Forecasting - Prophet",
                 "Forecasting - LSTM", "Forecasting - Hybrid", "Churn Prediction", "Inventory Optimization", "Tuning", "Drift Detection",
                 "Model Refinement", "MLOps - MLflow"]

COMPACT_KEYS = ["00_architecture_diagram", "01_03_monthly_revenue_orders", "01_15_pareto", "02_01_cleaning_waterfall", "03_06_pca_clusters",
                "03_08_cluster_profile_heatmap", "04_05_decomposition", "05_01_prophet_forecast", "08_02_all_models_test", "08_06_future_forecast",
                "09_05_roc_curves", "09_12_shap_summary", "10_08_backtest", "12_03_psi_simulated", "14_02_acceptance_scorecard", "11_04_default_vs_tuned"]


# ----------------------------------------------------------------------------------- helpers
def g(d, *keys, default="n/a"):
    for k in keys:
        if isinstance(d, dict) and k in d:
            d = d[k]
        else:
            return default
    return d


def fnum(x, nd=2, default="n/a"):
    try:
        return f"{float(x):,.{nd}f}"
    except Exception:
        return default


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for t, txt in (("begin", None), (None, instr), ("end", None)):
        if t:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), t)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = NAVY if level <= 2 else TEAL
        r.font.name = "Calibri"
    return h


def para(doc, text, bold=False, italic=False, size=10.5, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(it)
        r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(2)


def table(doc, header, rows, widths=None, font=9, head_fill="1F3A5F"):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold, r.font.size, r.font.color.rgb = True, Pt(font), RGBColor(255, 255, 255)
        shade(c, head_fill)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font)
            if ri % 2 == 1:
                shade(cells[i], "F1F5F9")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def df_rows(df, nd=2):
    return [[str(idx)] + [fnum(v, nd) if isinstance(v, (int, float)) else str(v) for v in row] for idx, row in zip(df.index, df.values.tolist())]


def jpg_bytes(path, max_w=1500, q=86):
    from PIL import Image
    im = Image.open(path).convert("RGB")
    if im.width > max_w:
        im = im.resize((max_w, int(im.height * max_w / im.width)))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=q, optimize=True)
    buf.seek(0)
    return buf


class Fig:
    n = 0


def figure(doc, key, idx, width=6.3, caption=True):
    meta = idx.get(key)
    if not meta or not Path(meta["file"]).exists():
        return False
    Fig.n += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(jpg_bytes(meta["file"]), width=Inches(width))
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(f"Figure {Fig.n}: {meta['title']}")
        r.bold, r.font.size = True, Pt(9)
        if meta.get("note"):
            c2 = doc.add_paragraph()
            c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = c2.add_run(meta["note"])
            r2.italic, r2.font.size, r2.font.color.rgb = True, Pt(8.5), GREY
            c2.paragraph_format.space_after = Pt(8)
    return True


def figure_grid(doc, keys, idx, per_row=2, width=3.15):
    keys = [k for k in keys if k in idx and Path(idx[k]["file"]).exists()]
    for i in range(0, len(keys), per_row):
        t = doc.add_table(rows=2, cols=per_row)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, k in enumerate(keys[i:i + per_row]):
            Fig.n += 1
            cell = t.rows[0].cells[j]
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.paragraphs[0].add_run().add_picture(jpg_bytes(idx[k]["file"], 1100), width=Inches(width))
            cap = t.rows[1].cells[j]
            cap.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.paragraphs[0].add_run(f"Fig. {Fig.n}: {idx[k]['title']}")
            r.bold, r.font.size = True, Pt(8)
            if idx[k].get("note"):
                r2 = cap.add_paragraph().add_run(idx[k]["note"][:170])
                r2.italic, r2.font.size, r2.font.color.rgb = True, Pt(7.5), GREY
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ----------------------------------------------------------------------------------- main
def build_report(out_path: str, compact: bool = False) -> str:
    Fig.n = 0
    M = ru.get_metrics()
    idx = json.loads(ru.CHART_INDEX.read_text()) if ru.CHART_INDEX.exists() else {}
    data, fc, hy, ch = M.get("data", {}), M.get("forecast", {}), M.get("hybrid", {}), M.get("churn", {})
    inv, seg, tun, drift, clean_, ts = M.get("inventory", {}), M.get("segmentation", {}), M.get("tuning", {}), M.get("drift", {}), M.get("cleaning", {}), M.get("timeseries", {})
    acc = M.get("acceptance", {})

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)   # A4
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)

    hp = sec.header.paragraphs[0]
    hp.text = f"RetailPulse  |  Project made by {AUTHOR}  |  Zidio Development - Data Science & Analytics"
    hp.runs[0].font.size, hp.runs[0].font.color.rgb = Pt(8), GREY
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    r.font.size = Pt(8)
    add_field(fp, "PAGE")

    # ---------------- cover
    if "00_project_banner" in idx and Path(idx["00_project_banner"]["file"]).exists():
        p = doc.add_paragraph()
        p.add_run().add_picture(jpg_bytes(idx["00_project_banner"]["file"]), width=Inches(6.6))
    para(doc, "RetailPulse", bold=True, size=34, color=NAVY, space_after=0)
    para(doc, "AI-Powered Customer Analytics & Demand Forecasting Platform", bold=True, size=17, color=TEAL, space_after=4)
    para(doc, "Predictive Demand  •  Customer Segmentation  •  Churn Analysis  •  Inventory Optimization", size=11, color=GREY, space_after=18)
    table(doc, ["Item", "Details"], [
        ["Project made by", AUTHOR],
        ["Prepared for", "Zidio Development - Data Science & Analytics Domain"],
        ["Document", "Project Report (compact Zidio submission edition)" if compact else "Complete Project Report (every chart explained)"],
        ["Date / Version", "March 2026  |  2.0 - Industry Edition"],
        ["Dataset used", str(g(data, "source"))],
        ["Data period", f"{g(data, 'date_min')}  to  {g(data, 'date_max')}"],
    ], widths=[1.8, 4.8], font=10)
    para(doc, "Tagline: From raw transactions to forecasts, retention actions and replenishment plans - in one reproducible, monitored pipeline.", italic=True, color=GREY)
    page_break(doc)

    # ---------------- project summary
    heading(doc, "Project Summary", 1)
    para(doc, "RetailPulse is an end-to-end Data Science & Analytics platform that helps retail businesses make smarter decisions. "
              "It ingests sales, customer and inventory data and delivers four solutions: (1) demand forecasting for the next 30 days with a Prophet + LSTM hybrid, "
              "(2) customer segmentation with RFM, K-Means and DBSCAN, (3) churn prediction with XGBoost and SHAP explainability, and (4) inventory optimisation "
              "(safety stock, reorder point, EOQ and recommended order quantities). Results are delivered through a 5-page Streamlit dashboard, a FastAPI service and a "
              "12-page Power BI report, and the system follows MLOps practice: MLflow tracking, PSI/KS drift detection, an Airflow retraining DAG, Docker, Kubernetes and GitHub Actions CI/CD.")
    heading(doc, "Key results of this run (computed, not hard-coded)", 2)
    h = hy if hy else {}
    cm = ch if ch else {}
    table(doc, ["Area", "Metric", "Result", "Target"], [
        ["Data", "Clean rows / customers / products", f"{fnum(g(data, 'final_rows', default=None), 0)} / {fnum(g(data, 'customers', default=None), 0)} / {fnum(g(data, 'products', default=None), 0)}", "-"],
        ["Data", "Total revenue", fnum(g(data, "total_revenue", default=None), 0), "-"],
        ["Segmentation", "K-Means clusters (silhouette)", f"{g(seg, 'k_chosen')} ({fnum(g(seg, 'silhouette_kmeans', default=None), 3)})", "6-8 segments"],
        ["Forecasting", "Hybrid MAPE weekly totals", f"{fnum(g(h, 'MAPE_weekly_%', default=None), 1)} %", "<= 12 %"],
        ["Forecasting", "Hybrid MAPE daily / WAPE", f"{fnum(g(h, 'MAPE_daily_%', default=None), 1)} % / {fnum(g(h, 'WAPE_%', default=None), 1)} %", "informational"],
        ["Churn", "XGBoost AUC-ROC", fnum(g(cm, "AUC-ROC", default=None), 3), ">= 0.88"],
        ["Churn", "Precision @ top 20 %", fnum(g(cm, "Precision@top20%", default=None), 3), ">= 0.75"],
        ["Inventory", "Cost reduction / lost-sales reduction (back-test)", f"{fnum(g(inv, 'cost_reduction_pct', default=None), 1)} % / {fnum(g(inv, 'lost_sales_reduction_pct', default=None), 1)} %", "25-40 %"],
        ["MLOps", "Drift simulation triggers retrain", str(g(drift, "retrain_triggered_on_simulation")), "Yes"],
    ], widths=[1.1, 2.6, 1.8, 1.1], font=9)
    para(doc, "Honesty note: targets are taken from the project brief. Where a target is not met on the dataset used, the report states it openly "
              "(see the acceptance scorecard). Daily MAPE is high for lumpy wholesale demand, therefore weekly-total MAPE is reported as the planning metric.", italic=True, size=9.5, color=GREY)
    page_break(doc)

    # ---------------- contents
    heading(doc, "Contents", 1)
    toc = ["Project Summary", "1. Introduction & Business Case", "2. Project Overview", "3. Key Features (F-01 to F-06)", "4. Technology Stack", "5. Architecture & Workflow",
           "6. Dataset, Cleaning & New Columns", "7. 28-Day Execution Timeline"] + [f"{8 + i}. {SECTION_TEXT[s][0]}" for i, s in enumerate([s for s in SECTION_ORDER if any(v['section'] == s for v in idx.values())])]
    n_ = 8 + len([s for s in SECTION_ORDER if any(v['section'] == s for v in idx.values())])
    toc += [f"{n_}. Acceptance Scorecard", f"{n_ + 1}. Deployment & Operations", f"{n_ + 2}. Power BI Dashboards (12 pages)", f"{n_ + 3}. Challenges & How They Were Solved",
            f"{n_ + 4}. Personal Reflection", "Appendix A - Chart index", "Appendix B - Repository structure & how to run"]
    bullets(doc, toc)
    page_break(doc)

    # ---------------- 1-2
    heading(doc, "1. Introduction & Business Case", 1)
    para(doc, "Retailers lose billions because of poor demand forecasting and stock mismanagement: shelves run empty when demand is under-estimated and capital is locked in slow stock when it is over-estimated. "
              "At the same time, acquiring a new customer costs far more than keeping an existing one, so identifying customers who are about to leave is valuable.")
    para(doc, "Mission: build an end-to-end data science platform that ingests sales, customer and inventory data and delivers accurate demand forecasts, customer segmentation, churn prediction and inventory optimisation recommendations.")
    bullets(doc, ["Reduce stock-outs through accurate demand forecasting", "Increase revenue through better inventory decisions", "Improve retention by identifying at-risk customers early",
                  "Process 10M+ transactions per month with daily batch jobs under 5 minutes (design target; large files are sampled by customer in Colab)"])
    heading(doc, "2. Project Overview", 1)
    table(doc, ["Topic", "Description"], [
        ["Vision", "One analytics platform that turns retail transactions into forecasts, customer actions and replenishment plans."],
        ["Target users", "Supermarket chains, fashion retailers and e-commerce companies: merchandisers, supply-chain planners, CRM / marketing teams, management."],
        ["Business value", "Fewer stock-outs, lower holding cost, targeted retention campaigns, one source of truth dashboard."],
        ["Non-functional goals", "MAPE <= 12 % (weekly totals), daily batch < 5 min, scalable to 10M+ transactions / month, full MLflow tracking and drift detection."],
    ], widths=[1.4, 5.2])
    heading(doc, "3. Key Features", 1)
    mp = lambda k: ("Met" if g(acc, k, "met", default=None) else ("Not met" if g(acc, k, "met", default=None) is False else "n/a"))
    table(doc, ["ID", "Feature", "Description", "Acceptance criteria", "Status"], [
        ["F-01", "Data ingestion & cleaning", "Kaggle / upload / demo loaders, ETL, validation", "Automated pipeline + quality checks", mp("F-01")],
        ["F-02", "Customer segmentation", "RFM + K-Means / DBSCAN", "6-8 segments with interpretation", mp("F-02")],
        ["F-03", "Demand forecasting", "Prophet + LSTM hybrid, 30 days", "MAPE <= 12 %", mp("F-03")],
        ["F-04", "Churn prediction", "XGBoost + SHAP", "AUC >= 0.88, precision@20 % >= 0.75", f"{mp('F-04a')} / {mp('F-04b')}"],
        ["F-05", "Inventory optimisation", "Safety stock, ROP, EOQ, order qty", "Reduce over/under-stock 25-40 %", mp("F-05")],
        ["F-06", "Interactive dashboard", "Streamlit (5 pages) + Power BI (12 pages)", "Insights, what-if, exports", "Delivered"],
    ], widths=[0.5, 1.3, 1.9, 1.9, 0.9], font=8.5)
    heading(doc, "4. Technology Stack", 1)
    table(doc, ["Category", "Technology", "Rationale / alternatives"], [
        ["Language", "Python 3.11", "Data-science ecosystem"], ["Data processing", "Pandas, NumPy, Scikit-learn", "Core manipulation and ML"],
        ["Forecasting", "Prophet + LSTM (PyTorch)", "Hybrid: calendar effects + non-linear dynamics"], ["Churn / explainability", "XGBoost, SHAP", "Best tabular accuracy + transparent drivers"],
        ["Tuning", "Optuna", "Bayesian search, faster than grid search"], ["Dashboard", "Streamlit, Plotly, Power BI", "Fast interactive analytics / business reporting"],
        ["Serving", "FastAPI", "Typed, fast REST endpoints"], ["Experiment tracking", "MLflow (SQLite)", "Versioning and reproducibility, no paid service"],
        ["Monitoring", "PSI / KS, Evidently AI, Prometheus + Grafana", "Drift and performance monitoring"], ["Orchestration", "Airflow DAG, GitHub Actions", "Retraining and CI/CD"],
        ["Deployment", "Docker, Kubernetes, Streamlit Cloud", "Consistent, scalable deployment"]], widths=[1.4, 2.3, 2.9])

    heading(doc, "5. Architecture & Workflow", 1)
    para(doc, "RetailPulse has five layers: Data -> Feature Engineering (RFM, lags, rolling averages) -> Models (segmentation, forecasting, churn, inventory, MLflow registry) -> Serving (FastAPI, Streamlit, Power BI) -> Monitoring (drift, performance tracking) with a retraining feedback loop.")
    figure(doc, "00_architecture_diagram", idx, 6.4)
    figure(doc, "00_workflow", idx, 6.4)

    heading(doc, "6. Dataset, Cleaning & New Columns", 1)
    para(doc, f"Data source used for this run: {g(data, 'source')}. Raw rows: {fnum(g(data, 'raw_rows', default=None), 0)}; clean sales rows: {fnum(g(data, 'final_rows', default=None), 0)} "
              f"({fnum(g(clean_, 'retention_pct', default=None), 1)} % retained). Customers: {fnum(g(data, 'customers', default=None), 0)}, products: {fnum(g(data, 'products', default=None), 0)}, countries: {fnum(g(data, 'countries', default=None), 0)}.")
    table(doc, ["Cleaning step", "Rows"], [[k.replace("_", " "), fnum(g(data, k, default=None), 0)] for k in
                                           ["raw_rows", "duplicates_removed", "returns_separated", "invalid_qty_or_price_removed", "non_product_codes_removed", "extreme_outliers_removed", "rows_without_customer_id", "final_rows"]], widths=[4, 2.2])
    para(doc, "Two columns were added to the cleaned CSV / Power BI file:", bold=True)
    bullets(doc, ["invoice_no - numeric part of the invoice id. Effect: a clean integer key for distinct-order counts, joins with returns and Power BI relationships (removes prefixes such as 'C' or 'A').",
                  "order_value - total revenue of the whole invoice repeated on each line. Effect: average-order-value, basket-value bands and order-level filters work without re-grouping, and visuals stay fast."])
    page_break(doc)

    heading(doc, "7. 28-Day Execution Timeline", 1)
    table(doc, ["Week", "Days", "Deliverables", "Notebook(s)"], [
        ["1 - Exploration & preparation", "1-7", "EDA, cleaning, RFM segmentation, time-series prep, Prophet, LSTM, MLflow baselines", "01-07"],
        ["2 - Advanced modelling", "8-14", "Hybrid forecaster, churn + SHAP, inventory, Optuna, drift detection, retraining pipeline", "08-14"],
        ["3 - Dashboard & analytics", "15-21", "Streamlit 5 pages, what-if, alerts, CSV/PDF export, Power BI (12 pages)", "15 + app/"],
        ["4 - Deployment & polish", "22-28", "Docker, Kubernetes, GitHub Actions, cloud deploy, monitoring, load test, report, demo video", "Dockerfile, k8s/, .github/"]],
        widths=[1.7, 0.6, 3.5, 0.9], font=8.5)

    # ---------------- technical sections
    n = 8
    present = [s for s in SECTION_ORDER if any(v["section"] == s for v in idx.values())]
    for s in present:
        title, text = SECTION_TEXT.get(s, (s, ""))
        heading(doc, f"{n}. {title}", 1)
        para(doc, text)
        if s in ("Forecasting - LSTM", "Forecasting - Hybrid") and "fallback" in str(fc.get("engines", "")).lower():
            para(doc, "Note for this run: PyTorch was not installed in the environment that produced these numbers, so a gradient-boosting lag model stood in for the LSTM "
                      "(the engine is printed in the notebook). Re-run the notebooks in Google Colab, where PyTorch is pre-installed, to obtain the true LSTM and ensemble results.", italic=True, size=9.5, color=GREY)
        if s == "Forecasting - Hybrid" and fc.get("table"):
            tb = fc["table"]
            cols = ["MAPE_daily_%", "MAPE_weekly_%", "WAPE_%", "sMAPE_%", "RMSE", "MAE"]
            table(doc, ["Model"] + cols, [[m] + [fnum(v.get(c), 2) for c in cols] for m, v in tb.items()], font=8)
            para(doc, f"Blend weight (Prophet) = {fnum(fc.get('blend_weight_prophet'), 2)}; engines: {fc.get('engines')}.", italic=True, size=9, color=GREY)
        if s == "Churn Prediction" and ch.get("table"):
            cols = ["AUC-ROC", "PR-AUC", "F1", "Precision@top20%", "Recall@top20%", "Brier"]
            table(doc, ["Model"] + cols, [[m] + [fnum(v.get(c), 3) for c in cols] for m, v in ch["table"].items()], font=8)
        if s == "Customer Segmentation" and seg.get("segments"):
            table(doc, ["Segment", "Customers", "Revenue share %"], [[k, fnum(v, 0), fnum(g(seg, "revenue_share_pct", k, default=None), 1)] for k, v in seg["segments"].items()], widths=[3, 1.5, 1.5])
        if s == "Inventory Optimization" and inv:
            table(doc, ["Metric", "Value"], [["Lost-sales reduction (back-test)", f"{fnum(inv.get('lost_sales_reduction_pct'), 1)} %"], ["Total cost reduction", f"{fnum(inv.get('cost_reduction_pct'), 1)} %"],
                                              ["Inventory level change", f"{fnum(inv.get('inventory_change_pct'), 1)} %"], ["Fill rate naive -> RetailPulse", f"{fnum(inv.get('fill_rate_naive_pct'), 1)} % -> {fnum(inv.get('fill_rate_opt_pct'), 1)} %"],
                                              ["SKUs: reorder now / overstock", f"{inv.get('reorder_now', 'n/a')} / {inv.get('overstock', 'n/a')}"]], widths=[3.5, 2.5])
        if s == "Tuning" and tun:
            table(doc, ["Item", "Value"], [["XGBoost default test AUC", fnum(tun.get("xgb_default_test_auc"), 4)], ["XGBoost tuned test AUC", fnum(tun.get("xgb_tuned_test_auc"), 4)],
                                           ["Prophet best validation WAPE", f"{fnum(tun.get('prophet_best_val_wape'), 2)} %"]], widths=[3.5, 2.5])
        keys = [k for k, v in idx.items() if v["section"] == s]
        if compact:
            keys = [k for k in keys if k in COMPACT_KEYS]
            figure_grid(doc, keys, idx, 2, 3.15)
        else:
            for k in keys:
                figure(doc, k, idx, 6.0)
        n += 1

    # ---------------- acceptance
    heading(doc, f"{n}. Acceptance Scorecard", 1)
    if acc:
        table(doc, ["ID", "Requirement", "Result", "Status"], [[k, v.get("requirement"), v.get("result"), {True: "MET", False: "NOT MET", None: "NOT RUN"}[v.get("met")]] for k, v in acc.items()], widths=[0.6, 2.5, 2.6, 0.9], font=8.5)
    else:
        para(doc, "Run notebook 14 to generate the acceptance scorecard.")
    figure(doc, "14_02_acceptance_scorecard", idx, 5.6) if not compact else None
    n += 1

    heading(doc, f"{n}. Deployment & Operations", 1)
    bullets(doc, ["Platform: Streamlit Community Cloud (free, HTTPS) or Hugging Face Spaces; container image for AWS / GCP.",
                  "Docker: multi-stage image (builder + slim runtime) running Streamlit on port 8501 with a health check.",
                  "Kubernetes: Deployment (2 replicas, liveness / readiness probes), Service and Ingress manifests in k8s/.",
                  "CI/CD: GitHub Actions runs ruff, pytest and a smoke test of the pipeline on every push / pull request, then builds the Docker image.",
                  "Serving: FastAPI endpoints /health, /forecast, /churn/predict and /segment.",
                  "Monitoring: PSI / KS drift tables, rolling WAPE monitor, Evidently HTML report, weekly Airflow DAG that retrains when drift is detected; Prometheus + Grafana recommended for production metrics.",
                  "Security: customer data anonymisation, role-based access in the dashboard (planned), JWT-secured API endpoints (recommended), audit logging of sensitive operations."])
    n += 1
    heading(doc, f"{n}. Power BI Dashboards (12 pages)", 1)
    para(doc, "A Power BI-ready file (powerbi_ready.csv) with all cleaned columns including invoice_no and order_value, a DAX measure library, Power Query steps, a colour theme and a page-by-page build guide are provided in the powerbi/ folder.")
    table(doc, ["Page", "Dashboard", "Key visuals"], [
        ["1", "Executive Summary", "KPI cards, revenue trend, top country, top product"], ["2", "Sales Trend Analysis", "Line + MoM / YoY, moving average"],
        ["3", "Product Performance", "Top / bottom products, category treemap, Pareto"], ["4", "Customer Analytics", "Customers, new vs returning, RFM segments"],
        ["5", "Region / Country Sales", "Map, country ranking"], ["6", "Monthly & Seasonal Sales", "Month x weekday matrix, seasonality"],
        ["7", "Customer Behaviour", "Frequency, basket size, cohort matrix"], ["8", "Profit & Revenue Analysis", "Revenue, estimated margin, price bands"],
        ["9", "Inventory Risk", "Reorder-now SKUs, days of cover, ABC"], ["10", "Order & Transaction", "Orders (distinct invoice_no), order_value bands"],
        ["11", "Advanced Analytics", "Forecast vs actual, churn risk, decomposition"], ["12", "Interactive Filters", "Slicers, drill-through, bookmarks"]], widths=[0.5, 2.0, 4.1], font=8.5)
    n += 1
    heading(doc, f"{n}. Challenges & How They Were Solved", 1)
    table(doc, ["Challenge", "Solution"], [
        ["Non-stationary, lumpy daily demand with closed weekdays", "Decomposition, ADF/KPSS, closed-day detection, weekly-total MAPE as planning metric, hybrid model"],
        ["Churn label leakage", "Features built only from history before a cut-off, label from the following 90 days, chronological logic"],
        ["Accuracy vs interpretability", "XGBoost with SHAP global and per-customer explanations"],
        ["No inventory table in public data", "Stock on hand simulated and clearly labelled; policy evaluated with a real-demand back-test"],
        ["Huge files (10M rows) in Colab", "Customer-level hash sampling keeps whole customer histories while limiting memory"],
        ["Library API changes (Evidently, Great Expectations)", "Native PSI/KS and validation checks plus optional Evidently with version fallback"]], widths=[2.7, 3.9], font=8.5)
    n += 1
    heading(doc, f"{n}. Personal Reflection", 1)
    para(doc, f"By {AUTHOR}. (Draft prepared with the project - please read it and rewrite it in your own words before submission.)", italic=True, size=9.5, color=GREY)
    para(doc, "Key learnings: a good forecast starts with understanding the calendar - closed days, weekly rhythm and holiday peaks mattered more than model complexity. "
              "Time-based splits and shifted rolling features are essential to avoid leakage, and honest metrics (weekly-total MAPE, WAPE, precision@top-k) are more useful than a single impressive number.")
    para(doc, "Industry best practices applied: reproducible environment, MLflow tracking, validation checks, drift monitoring with an automatic retrain trigger, explainable models and CI/CD.")
    para(doc, "Future roadmap: hierarchical (SKU x store) forecasting, promotion and price features, real inventory feed, uplift modelling for retention offers, feature store, streaming ingestion and automated A/B evaluation of campaigns.")
    page_break(doc)

    heading(doc, "Appendix A - Chart index", 1)
    rows = [[i + 1, v["title"], v["section"], k] for i, (k, v) in enumerate(sorted(idx.items()))]
    table(doc, ["#", "Chart", "Section", "File"], rows, widths=[0.4, 3.0, 1.7, 1.6], font=7.5)
    heading(doc, "Appendix B - Repository structure & how to run", 1)
    para(doc, "RetailPulse/ - notebooks/ (00 MAIN + 01-15), src/ (retailpulse_utils.py, retailpulse_report.py), app/ (Streamlit), api/ (FastAPI), data/ (sample + processed), powerbi/, k8s/, airflow/, scripts/, tests/, .github/workflows/, Dockerfile, requirements.txt, README.md.", size=10)
    bullets(doc, ["Colab: open notebooks/00_RetailPulse_MAIN.ipynb -> Runtime -> Run all (dataset loads from Kaggle, upload or demo).", "Local: pip install -r requirements.txt; streamlit run app/streamlit_app.py.",
                  "Tests: pytest -q.   Retrain: python scripts/retrain_pipeline.py --force."])
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else str(ru.REPORTS / "RetailPulse_Project_Report_FULL.docx")
    print(build_report(out, compact="--compact" in sys.argv))
